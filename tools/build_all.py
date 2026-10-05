"""Rebuild the complete lecture package from editable sources, offline after pip.

Run from any directory: python /path/to/repo/tools/build_all.py
The separate closing-demo application is deliberately not generated.
"""
from __future__ import annotations
from pathlib import Path
import contextlib,hashlib,io,json,os,platform,sys,unittest,zipfile
from importlib.metadata import version
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))

def build():
 os.chdir(ROOT)
 for d in ['graphs','web','slides','handouts','checkpoints','checkpoints/solutions','notebooks','results','fallback','dist']: (ROOT/d).mkdir(parents=True,exist_ok=True)
 from tools import graphics,make_checkpoints,build_web,make_slides,make_handouts,make_fallback
 from tools.benchmark import run_benchmarks
 print('1/8 Shared diagrams and fixtures');graphics.build(ROOT)
 print('2/8 Checkpoints and computed notebook outputs');make_checkpoints.build(ROOT)
 print('3/8 Actual source traces and self-contained HTML');data=build_web.build(ROOT)
 print('4/8 Measured comparisons');rows=run_benchmarks(9,ROOT/'results/comparison.csv')
 (ROOT/'results/benchmark_console.txt').write_text('\n'.join(f"{r['case']:36} {r['algorithm']:16} {r['status']:20} cost={r['returned_cost']} expansions={r['expansions']} median_ms={r['median_ms']}" for r in rows)+'\n')
 print('5/8 PowerPoint, slide PDF, and speaker notes');make_slides.build(ROOT)
 print('6/8 Printable guides and fallback panels');make_handouts.build(ROOT);make_fallback.build(ROOT,data)
 print('7/8 Tests and artifact validation')
 stream=io.StringIO();suite=unittest.defaultTestLoader.discover(str(ROOT/'tests'));r=unittest.TextTestRunner(stream=stream,verbosity=2).run(suite)
 (ROOT/'results/test_report.txt').write_text(stream.getvalue())
 if not r.wasSuccessful():raise RuntimeError(stream.getvalue())
 import fitz
 from pptx import Presentation
 p=Presentation(ROOT/'slides/dijkstra_astar.pptx')
 assert len(p.slides)==33
 assert all(len(s.notes_slide.notes_text_frame.text)>40 for s in p.slides)
 pages={str(f.relative_to(ROOT)):len(fitz.open(f)) for folder in ['slides','handouts','fallback'] for f in (ROOT/folder).glob('*.pdf')}
 assert pages['handouts/student_worksheet.pdf']==2,pages
 assert pages['slides/dijkstra_astar.pdf']==33,pages
 assert pages['fallback/fallback.pdf']==6,pages
 assert len(rows)==56 and all(r['status'] in ['found','budget_exhausted','not_run_size_limit'] for r in rows)
 report=dict(status='passed',unit_tests=r.testsRun,python=platform.python_version(),slides=33,core_slides=25,appendix_slides=8,pdf_pages=pages,benchmark_rows=len(rows),dependencies={k:version(k) for k in ['python-pptx','reportlab','Pillow','CairoSVG','PyMuPDF']},exclusions=['Closing-demo application','Browser smoke tests are optional and run separately'])
 (ROOT/'results/build_validation.json').write_text(json.dumps(report,indent=2)+'\n')
 print('8/8 Manifest and complete distribution')
 def include(f):
  p=f.relative_to(ROOT)
  return f.is_file() and not any(x in ['.git','.github','__pycache__','.venv','.pytest_cache','dist','publication','_build'] for x in p.parts) and str(p)!='package-ready.json' and f.suffix not in ['.pyc','.pyo']
 files=sorted(f for f in ROOT.rglob('*') if include(f))
 manifest=[dict(path=str(f.relative_to(ROOT)),bytes=f.stat().st_size,sha256=hashlib.sha256(f.read_bytes()).hexdigest()) for f in files if f.name!='artifact_manifest.json']
 mf=ROOT/'docs/artifact_manifest.json';mf.write_text(json.dumps(manifest,indent=2)+'\n')
 if mf not in files:files.append(mf)
 with zipfile.ZipFile(ROOT/'dist/dijkstra-astar-lecture.zip','w',zipfile.ZIP_DEFLATED,compresslevel=9) as z:
  for f in sorted(files):z.write(f,str(f.relative_to(ROOT)))
 print(json.dumps(report,indent=2));print('Complete: dist/dijkstra-astar-lecture.zip')
 return report

if __name__=='__main__':build()
