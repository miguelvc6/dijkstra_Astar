"""Editable native-shape PowerPoint + matching vector PDF from one layout.

No slide screenshots are used as PowerPoint backgrounds. Text and all cards,
tables, lines, and figures' captions are editable. Graph illustrations are
provided separately as editable SVG and regenerable Python fixtures.
"""
from __future__ import annotations
import importlib.util
import math
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from algorithms import dijkstra_lazy,dijkstra_eager,astar
from graphs.main_graph import GRAPH,HEURISTIC

W,H=13.333333,7.5
BG='F3F1EB'; PAPER='FFFEFB'; INK='17323D'; MUTED='63777D'; TEAL='087E79'; GOLD='D78033'; LINE='D7DFDA'
REF_URLS={
 'R1':'https://ir.cwi.nl/pub/9256/9256D.pdf',
 'R2':'https://docs.python.org/3/library/heapq.html',
 'R3':'https://inst.eecs.berkeley.edu/~cs188/textbook/search/informed.html',
 'R4':'https://stanford-cs221.github.io/spring2023-extra/modules/search/search2.pdf',
 'R5':'https://networkx.org/documentation/stable/reference/algorithms/generated/networkx.algorithms.simple_paths.all_simple_paths.html',
 'R6':'https://www.rfc-editor.org/rfc/rfc2328.html#section-16.1',
 'R7':'https://www3.cs.stonybrook.edu/~rezaul/papers/TR-07-54.pdf',
 'R8':'https://arxiv.org/abs/2504.17033v2'}


def fonts():
    root=Path('/usr/share/fonts/truetype/dejavu')
    if root.exists():
        for name,file in [('Lecture','DejaVuSans.ttf'),('LectureB','DejaVuSans-Bold.ttf'),('LectureMono','DejaVuSansMono.ttf')]:
            pdfmetrics.registerFont(TTFont(name,str(root/file)))
        return 'Lecture','LectureB','LectureMono'
    # ReportLab bundles Vera. This avoids distributing any font files ourselves.
    import reportlab
    root=Path(reportlab.__file__).parent/'fonts'
    for name,file in [('Lecture','Vera.ttf'),('LectureB','VeraBd.ttf'),('LectureMono','VeraMono.ttf')]:
        if not (root/file).exists():file='Vera.ttf'
        pdfmetrics.registerFont(TTFont(name,str(root/file)))
    return 'Lecture','LectureB','LectureMono'

REG,BOLD,MONO=fonts()


def wrap(text,width,fs,font):
    out=[]
    for para in text.split('\n'):
        if not para:out.append('');continue
        words=para.split(' ');line=''
        for word in words:
            trial=line+' '+word if line else word
            if line and pdfmetrics.stringWidth(trial,font,fs)>width:
                out.append(line);line=word
            else:line=trial
        out.append(line)
    return out


class Painter:
    def __init__(self,prs,pdf):self.prs=prs;self.pdf=pdf;self.slide=None
    def start(self):
        self.slide=self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slide.background.fill.solid();self.slide.background.fill.fore_color.rgb=RGBColor.from_string(BG)
        self.pdf.setFillColor(HexColor('#'+BG));self.pdf.rect(0,0,W*72,H*72,stroke=0,fill=1)
    def rect(self,x,y,w,h,fill=PAPER,stroke=None,r=0):
        typ=MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE
        s=self.slide.shapes.add_shape(typ,Inches(x),Inches(y),Inches(w),Inches(h))
        if r:s.adjustments[0]=0.06
        s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(fill)
        if stroke:s.line.color.rgb=RGBColor.from_string(stroke);s.line.width=Pt(.7)
        else:s.line.fill.background()
        self.pdf.setFillColor(HexColor('#'+fill))
        if stroke:self.pdf.setStrokeColor(HexColor('#'+stroke));self.pdf.setLineWidth(.7)
        if r:self.pdf.roundRect(x*72,(H-y-h)*72,w*72,h*72,r*72,stroke=bool(stroke),fill=1)
        else:self.pdf.rect(x*72,(H-y-h)*72,w*72,h*72,stroke=bool(stroke),fill=1)
    def text(self,x,y,w,h,text,fs=20,color=INK,bold=False,mono=False,align='left',min_fs=13):
        font=MONO if mono else BOLD if bold else REG
        # Code preserves indentation and explicit line breaks. Other text wraps.
        def lines_at(size):return text.split('\n') if mono else wrap(text,w*72,size,font)
        lines=lines_at(fs)
        while fs>min_fs and (len(lines)*fs*1.2>h*72 or any(pdfmetrics.stringWidth(z,font,fs)>w*72 for z in lines)):
            fs-=.3;lines=lines_at(fs)
        if len(lines)*fs*1.2>h*72+1:
            raise ValueError(f'Text does not fit: {text[:100]} ({fs:.1f}pt, {len(lines)} lines)')
        for z in lines:
            if pdfmetrics.stringWidth(z,font,fs)>w*72+1:raise ValueError('Text too wide: '+z)
        tf=self.slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame
        tf.clear();tf.word_wrap=False;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        tf.vertical_anchor=MSO_ANCHOR.TOP
        from pptx.enum.text import PP_ALIGN
        for i,z in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.text=z;p.font.name='Consolas' if mono else 'Arial';p.font.size=Pt(fs)
            p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
            p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=1.2
            p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
        self.pdf.setFont(font,fs);self.pdf.setFillColor(HexColor('#'+color))
        for i,z in enumerate(lines):
            yy=(H-y)*72-fs*.88-i*fs*1.2
            if align=='center':self.pdf.drawCentredString((x+w/2)*72,yy,z)
            elif align=='right':self.pdf.drawRightString((x+w)*72,yy,z)
            else:self.pdf.drawString(x*72,yy,z)
    def image(self,path,x,y,w,h):
        im=Image.open(path);iw,ih=im.size;scale=min(w/iw,h/ih);ww,hh=iw*scale,ih*scale
        xx=x+(w-ww)/2;yy=y+(h-hh)/2
        self.slide.shapes.add_picture(str(path),Inches(xx),Inches(yy),width=Inches(ww),height=Inches(hh))
        self.pdf.drawImage(str(path),xx*72,(H-yy-hh)*72,ww*72,hh*72,mask='auto')
    def done(self,notes):
        self.slide.notes_slide.notes_text_frame.text=notes
        self.pdf.showPage()


def table(p,x,y,width,headers,rows,widths=None,fs=18,row_h=.62):
    widths=widths or [1/len(headers)]*len(headers)
    for i,row in enumerate([headers]+rows):
        xx=x
        for j,(value,fraction) in enumerate(zip(row,widths)):
            ww=width*fraction
            p.rect(xx,y+i*row_h,ww,row_h,INK if i==0 else PAPER if i%2 else 'EAF0E9')
            p.text(xx+.13,y+i*row_h+.13,ww-.26,row_h-.18,str(value),fs=fs,
                   color=PAPER if i==0 else INK,bold=(i==0 or j==0),min_fs=12)
            xx+=ww


def frame(p,s,i):
    p.text(.65,.34,10,.3,s.get('kicker','PATHFINDING'),fs=11,color=TEAL,bold=True,min_fs=10)
    p.text(11,.34,1.66,.3,s['time']+(' min' if s['time']!='Appendix' else ''),fs=11,color=MUTED,align='right',min_fs=10)
    p.text(.65,.88,12.05,.85,s['title'],fs=32,bold=True,min_fs=28)
    if 'bottom' in s:
        p.rect(.65,6.43,12.02,.58,INK,r=.08)
        p.text(.88,6.55,11.55,.38,s['bottom'],fs=15.4,color=PAPER,min_fs=11.8)
    labels=['MODEL','ENUMERATE','FINALIZE','INFORM','TEST']
    for j,l in enumerate(labels):
        xx=.65+j*1.7
        p.rect(xx,7.15,1.42,.025,TEAL if j==s['phase'] else LINE)
        p.text(xx,7.22,1.42,.15,l,fs=7.5,color=TEAL if j==s['phase'] else MUTED,bold=True,min_fs=7)
    refs=' · '.join(s.get('refs',[]))
    p.text(9.45,7.19,2.45,.22,refs,fs=9,color=MUTED,align='right',min_fs=8)
    p.text(12.0,7.16,.65,.25,f'{i:02d}',fs=11,color=INK,align='right',min_fs=10)


def build(root):
    root=Path(root)
    spec=importlib.util.spec_from_file_location('slide_content',root/'slides/content.py')
    mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
    prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
    prs.core_properties.title='Dijkstra and A*: when can a route be finalized?'
    prs.core_properties.subject='90-minute PhD lecture: conjectures, counterexamples, proofs and code'
    prs.core_properties.author='Lecture teaching package'
    pdf=canvas.Canvas(str(root/'slides/dijkstra_astar.pdf'),pagesize=(W*72,H*72))
    pdf.setTitle(prs.core_properties.title)
    p=Painter(prs,pdf)
    rows=[];notes=['# Slide map and speaker notes\n\n25 core slides; the appendix is not part of the 90-minute schedule.\n']
    for i,s in enumerate(mod.SLIDES,1):
        p.start();kind=s['kind']
        if kind=='cover':
            p.rect(0,0,5.65,H,INK)
            p.text(.7,.48,4.4,.5,'A 90-MINUTE PhD LECTURE',fs=12,color='8ED9C6',bold=True)
            p.text(.7,1.55,4.6,1.7,'Dijkstra\nand A*',fs=48,color=PAPER,bold=True,min_fs=45)
            p.text(.7,3.55,4.35,1.8,'When can a route\nbe finalized?',fs=29,color=PAPER,min_fs=27)
            p.text(.7,6.25,4.25,.68,'Conjecture → counterexample → repair',fs=15,color='C9DED4',min_fs=13)
            p.image(root/'graphs/main_graph.png',5.8,1.1,7.25,4.9)
            p.text(6.15,6.32,6.2,.68,'Find a route. Then justify stopping.',fs=21,color=TEAL,bold=True)
        else:
            frame(p,s,i)
            if kind=='cards':
                for j,(title,body) in enumerate(s['cards']):
                    x=.65+j*4.1
                    p.rect(x,2.13,3.82,3.78,PAPER,r=.1)
                    p.text(x+.28,2.42,3.24,.55,title,fs=24,bold=True)
                    p.rect(x+.28,3.14,.48,.04,TEAL)
                    p.text(x+.28,3.45,3.24,2.06,body,fs=21,min_fs=18)
            elif kind=='graph':
                p.rect(.65,1.97,7.62,4.22,PAPER,r=.1)
                p.image(root/'graphs'/s['image'],.76,2.04,7.4,4.1)
                p.text(8.65,2.32,3.9,.85,s['side_title'],fs=22,bold=True,min_fs=19)
                p.rect(8.65,3.17,.55,.045,TEAL)
                p.text(8.65,3.48,3.9,2.55,s['side'],fs=22,min_fs=18)
            elif kind=='formula':
                p.rect(.65,2.03,12.02,1.0,INK,r=.08)
                p.text(.95,2.3,11.4,.58,s['formula'],fs=29,color=PAPER,bold=True,min_fs=22)
                for j,k in enumerate(['left','right']):
                    title,body=s[k];x=.9+j*6.0
                    p.text(x,3.53,5.25,.5,title,fs=22,color=TEAL,bold=True)
                    p.text(x,4.3,5.25,1.8,body,fs=23,min_fs=18)
            elif kind=='code':
                p.rect(.65,2.03,7.76,4.14,INK,r=.12)
                p.text(.96,2.42,7.14,3.38,s['code'],fs=19,color='E3F2E9',mono=True,min_fs=15)
                p.text(8.78,2.38,3.8,.74,s['side_title'],fs=22,bold=True,min_fs=19)
                p.rect(8.78,3.25,.55,.045,TEAL)
                p.text(8.78,3.57,3.8,2.45,s['side'],fs=22,min_fs=18)
            elif kind=='proof':
                for j,(title,body) in enumerate(s['steps']):
                    yy=2.03+j*1.36
                    p.rect(.65,yy,12.02,1.12,PAPER,r=.08)
                    p.text(.92,yy+.23,.5,.56,str(j+1),fs=28,color=TEAL,bold=True)
                    p.text(1.68,yy+.16,2.1,.76,title,fs=22,bold=True)
                    p.text(3.91,yy+.18,8.3,.79,body,fs=20,min_fs=16)
            elif kind=='complexity':
                table(p,.65,2.13,12.02,['Variant','Time','Extra memory'],
                      [['Lazy / binary heap','O(n + m log(m+1))','O(n+m)'],
                       ['Eager / indexed heap','O((n+m) log(n+1))','O(n)']],
                      [.35,.4,.25],fs=21,row_h=.85)
                p.text(.92,5.22,11.55,.78,'Same search rule. Different update costs, queue size and implementation constants.',fs=22,min_fs=20)
            elif kind=='comparison':
                result=[dijkstra_lazy(GRAPH,'S','T'),dijkstra_eager(GRAPH,'S','T'),astar(GRAPH,'S','T',HEURISTIC)]
                rr=[[name,r.distance,r.stats.expansions,r.stats.stale_pops,r.stats.peak_queue] for name,r in zip(['Dijkstra / lazy','Dijkstra / eager','A*'],result)]
                table(p,.65,2.1,12.02,['Main graph','Cost','Expansions','Stale skips','Peak queue'],rr,[.34,.12,.2,.17,.17],fs=20,row_h=.75)
                p.text(.92,5.46,11.4,.55,'Fewer expansions, smaller queues, and lower runtime are different claims.',fs=21,color=TEAL,bold=True,min_fs=18)
            elif kind=='exit':
                for j,q in enumerate(s['questions']):
                    yy=2.18+j*1.33
                    p.text(.78,yy,.68,.65,f'{j+1:02d}',fs=25,color=TEAL,bold=True)
                    p.text(1.75,yy,10.4,.95,q,fs=26,min_fs=23)
            elif kind=='pathcounts':
                def count_paths(n):return sum(math.factorial(n-2)//math.factorial(n-2-k) for k in range(n-1))
                table(p,1.4,2.1,10.5,['Vertices n','Complete simple s-to-t paths'],[[n,f'{count_paths(n):,}'] for n in [5,7,9,11]],[.3,.7],fs=22,row_h=.7)
            elif kind=='research':
                p.text(.9,2.25,11.5,1.0,s['formula'],fs=50,color=TEAL,bold=True,min_fs=40)
                p.text(.94,3.94,11.4,1.78,s['side'],fs=25,min_fs=21)
            elif kind=='references':
                for j,(ident,title,url) in enumerate(s['entries']):
                    yy=2.06+j*1.11
                    p.text(.74,yy,.75,.5,ident,fs=18,color=TEAL,bold=True)
                    p.text(1.64,yy,10.85,.47,title,fs=18,min_fs=15)
                    p.text(1.64,yy+.51,10.85,.4,url,fs=12.8,color=MUTED,min_fs=11)
            else:raise ValueError(kind)
        note=f"SLIDE {i} | {s['time']} min | {'CORE' if i<=25 else 'OPTIONAL APPENDIX'}\n\n{s['notes']}"
        if s.get('refs'):note+='\n\nSOURCES\n'+'\n'.join(f"[{r}] {REF_URLS[r]}" for r in s['refs'])
        p.done(note)
        rows.append(f"| {i} | {s['time']} | {s['title']} |")
        notes.append(f"## {i:02d}. {s['title']} ({s['time']})\n\n{s['notes']}\n")
    prs.save(root/'slides/dijkstra_astar.pptx');pdf.save()
    (root/'slides/speaker_notes.md').write_text('\n'.join(notes)+'\n')
    (root/'slides/slide_map.md').write_text('# Slide map\n\n25 core slides + 8 optional appendix slides.\n\n| Slide | Minutes | Title |\n|---|---|---|\n'+'\n'.join(rows)+'\n')
    print(f'Created {len(mod.SLIDES)} slides with speaker notes.')


if __name__=='__main__':build(Path(__file__).resolve().parents[1])
